import os
import csv
import calendar
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ATTENDANCE_FILE = "data/attendance/attendance.csv"
EXCEL_FILE = "data/attendance/attendance.xlsx"
ENROLLED_DIR = "data/enrolled_faces"


def generate_excel(month=None, year=None):

    now = datetime.now()

    if month is None:
        month = now.month

    if year is None:
        year = now.year

    month_name = calendar.month_name[month]

    attendance = {}

    if os.path.exists(ATTENDANCE_FILE):
        with open(ATTENDANCE_FILE,"r",newline="",encoding="utf-8") as file:

            reader = csv.DictReader(file)

            for row in reader:

                if row.get("Status") != "Present":
                    continue

                name = row.get("Name")
                date = row.get("Date")

                if not name or not date:
                    continue

                try:
                    attendance_date = datetime.strptime(date,"%Y-%m-%d").date()
                except ValueError:
                    continue

                if (attendance_date.month == month and attendance_date.year == year):
                    attendance.setdefault(name, set()).add(attendance_date.day)

    students = []

    if os.path.isdir(ENROLLED_DIR):

        for name in os.listdir(ENROLLED_DIR):

            student_dir = os.path.join(
                ENROLLED_DIR,
                name
            )

            if os.path.isdir(student_dir):
                students.append(name)

    # Include anyone present in CSV
    # even if their enrollment folder
    # is currently unavailable.

    for name in attendance:

        if name not in students:
            students.append(name)

    students.sort(key=str.lower)

    # -----------------------------
    # Create workbook
    # -----------------------------

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"

    days_in_month = calendar.monthrange(
        year,
        month
    )[1]

    # -----------------------------
    # Title
    # -----------------------------

    last_column = days_in_month + 2

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column
    )

    sheet.cell(
        row=1,
        column=1,
        value="ATTENDANCE SHEET"
    )

    sheet.cell(
        row=1,
        column=1
    ).font = Font(
        size=20,
        bold=True
    )

    sheet.cell(
        row=1,
        column=1
    ).alignment = Alignment(
        horizontal="center"
    )

    # -----------------------------
    # Month / Year
    # -----------------------------

    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=last_column
    )

    sheet.cell(
        row=2,
        column=1,
        value=f"MONTH/YEAR: {month_name.upper()} {year}"
    )

    sheet.cell(
        row=2,
        column=1
    ).font = Font(
        bold=True,
        size=12
    )

    # -----------------------------
    # Header
    # -----------------------------

    sheet.cell(row=4, column=1, value="No.")
    sheet.cell(row=4, column=2, value="Name")

    # Week grouping
    week_start_column = 3
    week_number = 1

    for day in range(1, days_in_month + 1):

        date = datetime(
            year,
            month,
            day
        )
        INCLUDE_SATURDAY = True
        INCLUDE_SUNDAY =True
        if date.weekday() == 5 and not INCLUDE_SATURDAY:
            continue

        if date.weekday() == 6 and not INCLUDE_SUNDAY:
            continue

        column = day + 2

        sheet.cell(
            row=4,
            column=column,
            value=day
        )

        sheet.cell(
            row=5,
            column=column,
            value=date.strftime("%a")
        )

    # -----------------------------
    # Week labels
    # -----------------------------

    current_week_columns = []

    for day in range(1, days_in_month + 1):

        date = datetime(
            year,
            month,
            day
        )

        if date.weekday() >= 5:
            continue

        current_week_columns.append(day + 2)

        # Friday closes the week
        if date.weekday() == 4 or day == days_in_month:

            start_col = current_week_columns[0]
            end_col = current_week_columns[-1]

            sheet.merge_cells(
                start_row=3,
                start_column=start_col,
                end_row=3,
                end_column=end_col
            )

            sheet.cell(
                row=3,
                column=start_col,
                value=f"Week {week_number}"
            )

            sheet.cell(
                row=3,
                column=start_col
            ).alignment = Alignment(
                horizontal="center"
            )

            week_number += 1
            current_week_columns = []

    # -----------------------------
    # Student rows
    # -----------------------------

    start_row = 6

    for index, student in enumerate(
        students,
        start=1
    ):

        row = start_row + index - 1

        sheet.cell(
            row=row,
            column=1,
            value=index
        )

        sheet.cell(
            row=row,
            column=2,
            value=student
        )

        present_days = attendance.get(
            student,
            set()
        )

        for day in range(1, days_in_month + 1):

            date = datetime(
                year,
                month,
                day
            )
            INCLUDE_SUNDAY =False
            if date.weekday() == 5 and not INCLUDE_SATURDAY:
                continue
    
            if date.weekday() == 6 and not INCLUDE_SUNDAY:
                continue

            column = day + 2

            if day in present_days:
                value = "P"
            elif date.date() <= now.date():
                value = "-"
            else:
                value = ""

            sheet.cell(
                row=row,
                column=column,
                value=value
            )

    # -----------------------------
    # Formatting
    # -----------------------------

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row in sheet.iter_rows(
        min_row=3,
        max_row=sheet.max_row,
        min_col=1,
        max_col=last_column
    ):

        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    for cell in sheet[3]:
        cell.font = Font(bold=True)

    for cell in sheet[4]:
        cell.font = Font(bold=True)

    for cell in sheet[5]:
        cell.font = Font(bold=True)

    sheet.column_dimensions["A"].width = 6
    sheet.column_dimensions["B"].width = 22

    for column in range(3, last_column + 1):
        sheet.column_dimensions[
            get_column_letter(column)
        ].width = 5

    sheet.freeze_panes = "C6"

    # -----------------------------
    # Save
    # -----------------------------

    os.makedirs(
        os.path.dirname(EXCEL_FILE),
        exist_ok=True
    )

    workbook.save(EXCEL_FILE)

    print(
        f"Excel attendance sheet created: {EXCEL_FILE}"
    )


if __name__ == "__main__":
    generate_excel()